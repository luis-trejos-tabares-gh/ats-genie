import unicodedata
from io import BytesIO

from docx import Document
from fpdf import FPDF

from app.models.schemas import ResumeSections

# Helvetica (FPDF core font) is latin-1 only. Map common Unicode punctuation
# so bullets and compound words are not replaced with "?".
_PDF_CHAR_MAP = {
    "\u2022": "-",  # bullet
    "\u00b7": "-",  # middle dot
    "\u2010": "-",  # hyphen
    "\u2011": "-",  # non-breaking hyphen
    "\u2012": "-",  # figure dash
    "\u2013": "-",  # en dash
    "\u2014": "-",  # em dash
    "\u2212": "-",  # minus
    "\u00a0": " ",
    "\u202f": " ",
    "\u2009": " ",
    "\u200b": "",
    "\u2018": "'",
    "\u2019": "'",
    "\u201c": '"',
    "\u201d": '"',
    "\u2026": "...",
}

HEADINGS = {
    "en": {
        "contact": "Contact",
        "summary": "Summary",
        "skills": "Skills",
        "experience": "Experience",
        "education": "Education",
        "extras": "Additional",
        "resume": "Resume",
        "empty": "No content provided.",
    },
    "es": {
        "contact": "Contacto",
        "summary": "Resumen",
        "skills": "Competencias",
        "experience": "Experiencia",
        "education": "Educación",
        "extras": "Adicional",
        "resume": "Currículum",
        "empty": "No se proporcionó contenido.",
    },
    "pt": {
        "contact": "Contato",
        "summary": "Resumo",
        "skills": "Competências",
        "experience": "Experiência",
        "education": "Formação",
        "extras": "Adicional",
        "resume": "Currículo",
        "empty": "Nenhum conteúdo fornecido.",
    },
    "de": {
        "contact": "Kontakt",
        "summary": "Zusammenfassung",
        "skills": "Kenntnisse",
        "experience": "Berufserfahrung",
        "education": "Ausbildung",
        "extras": "Weitere Angaben",
        "resume": "Lebenslauf",
        "empty": "Kein Inhalt angegeben.",
    },
}


def _labels(language: str | None) -> dict[str, str]:
    return HEADINGS.get(language or "en", HEADINGS["en"])


def _contact_line(sections: ResumeSections) -> str:
    contact = sections.contact
    parts = [contact.name, contact.email, contact.phone, contact.location, *contact.links]
    return " | ".join(part for part in parts if part)


def _body_blocks(sections: ResumeSections, language: str | None) -> list[tuple[str, list[str]]]:
    labels = _labels(language)
    blocks: list[tuple[str, list[str]]] = []
    contact = _contact_line(sections)
    if contact:
        blocks.append((labels["contact"], [contact]))
    if sections.summary:
        blocks.append((labels["summary"], [sections.summary]))
    if sections.skills:
        blocks.append((labels["skills"], [", ".join(sections.skills)]))
    if sections.experience:
        lines: list[str] = []
        for item in sections.experience:
            heading = " | ".join(part for part in (item.title, item.employer) if part)
            dates = " - ".join(part for part in (item.startDate, item.endDate) if part)
            if heading or dates:
                lines.append(" ".join(part for part in (heading, dates) if part))
            lines.extend(f"- {bullet}" for bullet in item.bullets)
        blocks.append((labels["experience"], lines or [""]))
    if sections.education:
        lines = []
        for item in sections.education:
            lines.append(
                " | ".join(part for part in (item.school, item.credential, item.startDate, item.endDate) if part)
            )
        blocks.append((labels["education"], [line for line in lines if line]))
    if sections.extras:
        blocks.append((labels["extras"], [sections.extras]))
    if not blocks:
        blocks.append((labels["resume"], [labels["empty"]]))
    return blocks


def render_docx(sections: ResumeSections, language: str | None = None) -> bytes:
    document = Document()
    core = document.core_properties
    core.author = ""
    core.last_modified_by = ""
    core.comments = ""

    for heading, lines in _body_blocks(sections, language):
        document.add_heading(heading, level=1)
        for line in lines:
            document.add_paragraph(line)

    buffer = BytesIO()
    document.save(buffer)
    data = buffer.getvalue()
    return data.replace(b"python-docx", b"ats-assistant")


def _pdf_text(value: str) -> str:
    chars: list[str] = []
    for ch in value:
        mapped = _PDF_CHAR_MAP.get(ch, ch)
        if mapped != ch:
            chars.append(mapped)
            continue
        try:
            ch.encode("latin-1")
        except UnicodeEncodeError:
            decomposed = unicodedata.normalize("NFKD", ch).encode("latin-1", "ignore").decode("latin-1")
            chars.append(decomposed or "-")
        else:
            chars.append(ch)
    return "".join(chars)


def render_pdf(sections: ResumeSections, language: str | None = None) -> bytes:
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.add_page()
    pdf.set_creator("")
    pdf.set_author("")
    pdf.set_title("Resume")
    usable_width = pdf.w - pdf.l_margin - pdf.r_margin

    for heading, lines in _body_blocks(sections, language):
        pdf.set_x(pdf.l_margin)
        pdf.set_font("Helvetica", "B", 13)
        pdf.multi_cell(usable_width, 8, _pdf_text(heading))
        pdf.set_font("Helvetica", "", 11)
        for line in lines:
            pdf.set_x(pdf.l_margin)
            pdf.multi_cell(usable_width, 6, _pdf_text(line))
        pdf.ln(3)

    output = pdf.output()
    return bytes(output) if isinstance(output, (bytes, bytearray)) else output.encode("latin-1")


def render(sections: ResumeSections, fmt: str, language: str | None = None) -> tuple[bytes, str, str]:
    if fmt == "pdf":
        return render_pdf(sections, language), "application/pdf", "resume.pdf"
    return render_docx(sections, language), (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    ), "resume.docx"
